from flask import Flask, render_template, request, jsonify
import os
import tempfile

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB max

# ── Shop data ────────────────────────────────────────────────────────────────
SHOPS = [
    {
        "id": 1,
        "name": "AlWatan Print House",
        "location": "Manama Souq, Block 304, Bahrain",
        "contact": "+973 1723 4455",
        "services": ["Printing", "Scanning", "Binding", "Lamination"],
        "bw_price": 0.050,
        "color_price": 0.200,
        "rating": 4.8,
        "reviews": 142,
        "open": "Sat–Thu 8am–9pm",
        "image_url": "https://images.unsplash.com/photo-1562654501-a0ccc0fc3fb1?w=600&q=80",
        "color": "#1a3a5c"
    },
    {
        "id": 2,
        "name": "Gulf Copy Center",
        "location": "Seef District, Road 2832, Bahrain",
        "contact": "+973 1758 9900",
        "services": ["Printing", "Photocopying", "ID Photos", "Posters"],
        "bw_price": 0.040,
        "color_price": 0.180,
        "rating": 4.5,
        "reviews": 98,
        "open": "Daily 7am–10pm",
        "image_url": "https://images.unsplash.com/photo-1497366216548-37526070297c?w=600&q=80",
        "color": "#2d5a27"
    },
    {
        "id": 3,
        "name": "Express Print Bahrain",
        "location": "Juffair Ave, Block 215, Bahrain",
        "contact": "+973 1733 1122",
        "services": ["Printing", "Binding", "Business Cards", "Banners"],
        "bw_price": 0.060,
        "color_price": 0.220,
        "rating": 4.7,
        "reviews": 211,
        "open": "Sat–Thu 9am–8pm",
        "image_url": "https://images.unsplash.com/photo-1504868584819-f8e8b4b6d7e3?w=600&q=80",
        "color": "#5c1a1a"
    },
    {
        "id": 4,
        "name": "Muharraq Print Studio",
        "location": "Muharraq Old Town, Road 12, Bahrain",
        "contact": "+973 1761 5577",
        "services": ["Printing", "Scanning", "Photo Printing", "Stickers"],
        "bw_price": 0.035,
        "color_price": 0.150,
        "rating": 4.3,
        "reviews": 74,
        "open": "Sat–Wed 8am–7pm",
        "image_url": "https://images.unsplash.com/photo-1519389950473-47ba0277781c?w=600&q=80",
        "color": "#4a2c6e"
    },
    {
        "id": 5,
        "name": "Isa Town Office Solutions",
        "location": "Isa Town Market, Block 801, Bahrain",
        "contact": "+973 1768 3344",
        "services": ["Printing", "Binding", "Lamination", "Rubber Stamps"],
        "bw_price": 0.045,
        "color_price": 0.190,
        "rating": 4.6,
        "reviews": 133,
        "open": "Daily 8am–9pm",
        "image_url": "https://images.unsplash.com/photo-1497215728101-856f4ea42174?w=600&q=80",
        "color": "#7a4a00"
    },
    {
        "id": 6,
        "name": "Riffa Digital Prints",
        "location": "Riffa Central, Road 556, Bahrain",
        "contact": "+973 1777 2288",
        "services": ["Digital Printing", "Scanning", "Blueprints", "Canvas"],
        "bw_price": 0.055,
        "color_price": 0.230,
        "rating": 4.9,
        "reviews": 187,
        "open": "Sat–Thu 8:30am–8:30pm",
        "image_url": "https://images.unsplash.com/photo-1551288049-bebda4e38f71?w=600&q=80",
        "color": "#0a4a4a"
    },
]

def get_shop(shop_id):
    return next((s for s in SHOPS if s["id"] == shop_id), None)

# ── Page counting ─────────────────────────────────────────────────────────────
def count_pages(filepath, filename):
    ext = filename.rsplit('.', 1)[-1].lower()
    if ext == 'pdf':
        try:
            import pypdf
            with open(filepath, 'rb') as f:
                return len(pypdf.PdfReader(f).pages)
        except ImportError:
            raise Exception("pypdf not installed.")
    elif ext == 'docx':
        try:
            from docx import Document
            doc = Document(filepath)
            word_count = sum(len(p.text.split()) for p in doc.paragraphs)
            return max(1, round(word_count / 250))
        except ImportError:
            raise Exception("python-docx not installed.")
    elif ext == 'txt':
        with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
            lines = f.readlines()
        return max(1, round(len(lines) / 50))
    elif ext == 'pptx':
        try:
            from pptx import Presentation
            return len(Presentation(filepath).slides)
        except ImportError:
            raise Exception("python-pptx not installed.")
    else:
        raise Exception(f"Unsupported file type: .{ext}")

# ── Routes ────────────────────────────────────────────────────────────────────
@app.route('/')
def index():
    return render_template('index.html', shops=SHOPS)

@app.route('/shop/<int:shop_id>')
def shop(shop_id):
    s = get_shop(shop_id)
    if not s:
        return "Shop not found", 404
    return render_template('shop.html', shop=s)

@app.route('/upload', methods=['POST'])
def upload():
    shop_id = request.form.get('shop_id', type=int)
    print_type = request.form.get('print_type', 'bw')
    shop = get_shop(shop_id) if shop_id else None

    if 'file' not in request.files:
        return jsonify({'error': 'No file provided'}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400

    filename = file.filename
    allowed = {'pdf', 'docx', 'txt', 'pptx'}
    ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
    if ext not in allowed:
        return jsonify({'error': f'Unsupported file type. Allowed: {", ".join(allowed)}'}), 400

    with tempfile.NamedTemporaryFile(delete=False, suffix=f'.{ext}') as tmp:
        file.save(tmp.name)
        tmp_path = tmp.name

    try:
        pages = count_pages(tmp_path, filename)
        if shop:
            price_per_page = shop['color_price'] if print_type == 'color' else shop['bw_price']
        else:
            price_per_page = 0.25
        price = round(pages * price_per_page, 3)
        size_kb = round(os.path.getsize(tmp_path) / 1024, 1)
        return jsonify({
            'filename': filename,
            'pages': pages,
            'price': price,
            'price_per_page': price_per_page,
            'print_type': print_type,
            'size_kb': size_kb,
            'ext': ext.upper()
        })
    except Exception as e:
        return jsonify({'error': str(e)}), 500
    finally:
        os.unlink(tmp_path)

if __name__ == '__main__':
    app.run(debug=True)
